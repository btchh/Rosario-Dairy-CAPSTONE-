"""Parse the 2023-2025 annual sales workbooks without changing the database."""

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re

from openpyxl import load_workbook


CENT = Decimal('0.01')
DATE_TEXT = re.compile(r'^\s*(\d{1,2})/(\d{1,2})/(\d{2,4})\s*$')


def money(value):
    if value is None or value == '':
        return None
    try:
        result = Decimal(str(value))
    except (ValueError, ArithmeticError):
        return None
    return result.quantize(CENT, rounding=ROUND_HALF_UP) if result.is_finite() else None


def clean_label(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


CUSTOMER_ALIASES = {
    'outlet': 'Outlet Sales',
    'outletsales': 'Outlet Sales',
    'skye': 'Skye Dairy',
    'moo and moore': 'Moo and More',
    'tri star batangas corp': 'Tri Star',
    'arnold': 'Arnold Berse',
    'lito': 'Lito Ayson',
    'jamie': 'Jamie Ortiz',
    'joseph c.': 'Joseph Caguimbal',
    'wtc-svs': 'WTC_SVS',
    'wtc-sj': 'WTC_SJ',
    'wtc-a': 'WTC_A',
    'azaad': 'AZAAD',
    'azad': 'AZAAD',
    'ouutlet sales': 'Outlet Sales',
    'daiy box sales': 'Dairy Box Sales',
    'moo an more': 'Moo and More',
    'arnold besre': 'Arnold Berse',
    'agrijoy': 'Agri Joy',
    'eat plenty food corp.': 'Eat Plenty Food Corp',
    'eat plenty fodd corp': 'Eat Plenty Food Corp',
    'eat plemty food corp': 'Eat Plenty Food Corp',
    'eat pelnty food corp': 'Eat Plenty Food Corp',
    'eat plenty food': 'Eat Plenty Food Corp',
    'wtc__svs': 'WTC_SVS',
    'wtc_ svs': 'WTC_SVS',
    'pcc -uplb': 'PCC-UPLB',
    'milk feeding las pinas': 'Milk Feeding-Las Pinas',
    'milk feeding-las pinas': 'Milk Feeding-Las Pinas',
    'milk feeding dswd': 'Milk Feeding-DSWD',
    'milk feedin-dswd': 'Milk Feeding-DSWD',
    'milk feeding-dswd': 'Milk Feeding-DSWD',
    'milk feeding-deped': 'Milk Feeding-DEPED',
    'milk feeding rizal': 'Milk Feeding-Rizal',
    'milk feeding-rizal': 'Milk Feeding-Rizal',
    'kadiwa-sta.rosa': 'Kadiwa-Sta. Rosa',
    'kadiwa sta.rosa': 'Kadiwa-Sta. Rosa',
    'kadiwa sta rosa': 'Kadiwa-Sta. Rosa',
    'kadiwa_lares': 'Kadiwa-Lares',
    'kadiwa- alaminos': 'Kadiwa-Alaminos',
    'kadiwa- lumban': 'Kadiwa-Lumban',
}

PRODUCT_ALIASES = {
    'raw milk 1000ml': 'Raw Milk 1L',
    'rawmilk 1000ml': 'Raw Milk 1L',
    'raw milk l': 'Raw Milk 1L',
    'raw milk 1l': 'Raw Milk 1L',
    'raw milk': 'Raw Milk 1L',
    'raw l': 'Raw Milk 1L',
    'rawl': 'Raw Milk 1L',
    'cara 1000ml': 'Cara 1L',
    'cara l': 'Cara 1L',
    'cara 300ml': 'Cara 300ml',
    'cara 300': 'Cara 300ml',
    'cara300': 'Cara 300ml',
    'ice candy': 'Ice Candy',
    'uns l': 'Uns L',
    'unsl': 'Uns L',
    'mozza': 'Mozza',
    'mozarella': 'Mozza',
    'chocoballs': 'Choco Balls',
    'conebite': 'Cone Bite',
    'pdx13': 'PDLX13',
    'pdx16': 'PDLX16',
    'pld chz': 'PDL CHZ',
    'atssarang suka': 'Atsarang Suka',
    'samaploc': 'Sampaloc',
}


def canonical_customer(raw):
    label = clean_label(raw)
    if not label:
        return None
    label = re.sub(r'\(?\s*cancelled\s*\)?', '', label, flags=re.I).strip(' -_/')
    label = re.sub(r'\(SI\d+\)', '', label, flags=re.I).strip()
    if not label:
        return None
    if label.lower().startswith('outlet sales'):
        return 'Outlet Sales'
    if label.lower().startswith(('walk-in', 'walk in')):
        return 'Walk-in'
    label = re.sub(r'/\s*walk-in\b', '', label, flags=re.I).strip()
    return CUSTOMER_ALIASES.get(label.casefold(), label)


def customer_key(raw):
    name = canonical_customer(raw)
    return name.casefold() if name else None


def canonical_product(raw):
    label = clean_label(raw)
    key = label.casefold()
    if key in PRODUCT_ALIASES:
        return PRODUCT_ALIASES[key]
    if re.fullmatch(r'[a-z]{1,5}\d{2,4}', key) or re.fullmatch(r'pd[lx ]+\d{2,3}', key):
        return label.upper()
    return label


def category_for_product(name):
    key = name.casefold()
    if key.startswith(('raw milk', 'raw300', 'fml', 'cml', 'fm', 'cm', 'tm', 'mm', 'pm', 'bb', 'sb', 'cara', 'melon', 'pandan', 'ube', 'low fat', 'toned milk', 'pasteurized milk')) or key in {'ml', 'm300', 'm200', 'mml', 'uns l', 'sdl', 'sd300', 'pyl', 'p300', 'p200', 'pml', 'ice', 'raw'}:
        return 'Milk & Dairy Drinks'
    if any(word in key for word in ('cheese', 'mozza', 'yogurt', 'yoghurt', 'paneer')) or key in {'wc', 'mog'}:
        return 'Cheese & Cultured Dairy'
    if any(word in key for word in ('ice candy', 'pastillas', 'pdl', 'polvoron', 'peanut', 'chips', 'choco', 'cone bite', 'butcheron', 'banana strips', 'barquillos', 'cornick', 'dried fruit')):
        return 'Snacks & Desserts'
    return 'Other Products'


def source_date(value, sheet_month):
    if isinstance(value, datetime):
        parsed = value.date()
        if parsed.month != sheet_month and parsed.day == sheet_month:
            return date(parsed.year, parsed.day, parsed.month), 'swapped_excel_date'
        return parsed, 'sheet_mismatch' if parsed.month != sheet_month else None
    if isinstance(value, date):
        return value, 'sheet_mismatch' if value.month != sheet_month else None
    match = DATE_TEXT.match(value.replace(' ', '')) if isinstance(value, str) else None
    if match:
        month, day, year = map(int, match.groups())
        corrected_year = year == 223
        if corrected_year:
            year = 2023
        elif year < 100:
            year += 2000
        try:
            parsed = date(year, month, day)
        except ValueError:
            return None, 'invalid_date'
        return parsed, 'corrected_truncated_year' if corrected_year else ('sheet_mismatch' if month != sheet_month else None)
    return None, 'invalid_date' if value is not None else None


@dataclass
class SourceLine:
    row: int
    product: str
    quantity: Decimal
    unit_price: Decimal | None
    amount: Decimal | None


@dataclass
class SourceTicket:
    workbook: str
    sheet: str
    row: int
    sale_date: date | None
    customer_label: str
    invoice_number: str
    lines: list[SourceLine] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    adjustments: list[Decimal] = field(default_factory=list)
    reported_total: Decimal | None = None
    duplicate_of: str | None = None

    @property
    def reference(self):
        return f'{self.workbook}:{self.sheet}:{self.row}'

    @property
    def cancelled(self):
        return 'cancelled' in self.customer_label.casefold()

    @property
    def amount(self):
        if self.reported_total is not None and any(line.amount is None for line in self.lines):
            return self.reported_total
        return sum((line.amount or Decimal('0.00') for line in self.lines), Decimal('0.00')) + sum(self.adjustments, Decimal('0.00'))


@dataclass
class ParseResult:
    tickets: list[SourceTicket] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)
    daily_checks: list[dict] = field(default_factory=list)
    monthly_checks: list[dict] = field(default_factory=list)

    def issue(self, kind, workbook, sheet, row, detail):
        self.issues.append({'kind': kind, 'source': f'{workbook}:{sheet}:{row}', 'detail': detail})

    def summary(self):
        sold = [ticket for ticket in self.tickets if not ticket.cancelled and not ticket.duplicate_of and ticket.lines]
        return {
            'tickets': len(sold),
            'cancelled_tickets': sum(ticket.cancelled for ticket in self.tickets),
            'duplicate_tickets': sum(bool(ticket.duplicate_of) for ticket in self.tickets),
            'empty_tickets': sum(not ticket.lines and not ticket.cancelled and not ticket.duplicate_of
                                 for ticket in self.tickets),
            'line_items': sum(len(ticket.lines) for ticket in sold),
            'discount_rows': sum(len(ticket.adjustments) for ticket in sold),
            'revenue': str(sum((ticket.amount for ticket in sold), Decimal('0.00'))),
            'years': dict(Counter(ticket.sale_date.year for ticket in sold if ticket.sale_date)),
            'customers': len({customer_key(ticket.customer_label) for ticket in self.tickets
                              if not ticket.duplicate_of and customer_key(ticket.customer_label)}),
            'products': len({canonical_product(line.product).casefold() for ticket in sold for line in ticket.lines}),
            'issues': dict(Counter(issue['kind'] for issue in self.issues)),
            'invoice_mismatches': sum(check['difference'] != '0.00' for check in self.daily_checks),
            'monthly_mismatches': sum(abs(Decimal(check['difference'])) > Decimal('0.05') for check in self.monthly_checks),
        }


def parse_workbooks(source_dir):
    result = ParseResult()
    paths = sorted(Path(source_dir).glob('SALES REPORT JAN-DEC *.xlsx'))
    if len(paths) != 3:
        raise ValueError(f'Expected three annual workbooks in {source_dir}; found {len(paths)}.')
    for path in paths:
        workbook_year = int(path.stem[-4:])
        workbook = load_workbook(path, read_only=True, data_only=True)
        if len(workbook.worksheets) != 12:
            raise ValueError(f'{path.name} must have 12 monthly sheets.')
        for month, sheet in enumerate(workbook.worksheets, 1):
            current = None
            active_date = None
            sheet_tickets = []
            monthly_reported = None
            for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
                padded = list(row) + [None] * max(0, 8 - len(row))
                raw_date, raw_customer, raw_invoice, raw_product, raw_quantity, raw_price, raw_amount, raw_daily = padded[:8]
                if isinstance(raw_date, str) and raw_date.strip().upper() == 'TOTAL':
                    monthly_reported = money(raw_amount)
                    continue
                parsed_date, date_issue = source_date(raw_date, month)
                if raw_date is not None and date_issue and str(raw_date).strip().upper() != 'TOTAL':
                    result.issue(date_issue, str(workbook_year), sheet.title.strip(), row_number, str(raw_date))
                if parsed_date:
                    active_date = parsed_date
                customer = clean_label(raw_customer)
                invoice = clean_label(raw_invoice)
                product = clean_label(raw_product)
                if not any((customer, invoice, product)):
                    continue
                repeated_invoice = invoice and current and invoice == current.invoice_number and not customer and active_date == current.sale_date
                new_ticket = current is None or (bool(invoice) and not repeated_invoice) or bool(customer) or (parsed_date is not None and current.sale_date != parsed_date)
                if new_ticket:
                    current = SourceTicket(str(workbook_year), sheet.title.strip(), row_number, active_date, customer, invoice)
                    result.tickets.append(current)
                    sheet_tickets.append(current)
                    if active_date is None:
                        result.issue('missing_date', current.workbook, current.sheet, row_number, 'No date context')
                    if not invoice:
                        result.issue('missing_invoice', current.workbook, current.sheet, row_number, customer or product)
                    if not customer:
                        result.issue('missing_customer', current.workbook, current.sheet, row_number, invoice or product)
                if product:
                    if product in {'200', '300'} and current.lines:
                        size_codes = {'FML': 'FM', 'CML': 'CM', 'SBL': 'SB', 'BBL': 'BB', 'ML': 'M'}
                        base = next((line.product.upper().replace(' ', '') for line in reversed(current.lines)
                                     if line.product.upper().replace(' ', '') in size_codes), None)
                        if base:
                            product = size_codes[base] + product
                        else:
                            result.issue('unresolved_size_label', current.workbook, current.sheet, row_number, product)
                    if re.fullmatch(r'discoun\s*t', product, flags=re.I):
                        adjustment = money(raw_amount)
                        if adjustment is not None:
                            current.adjustments.append(adjustment)
                        else:
                            result.issue('missing_discount_amount', current.workbook, current.sheet, row_number, product)
                        if isinstance(raw_daily, (int, float, Decimal)):
                            current.reported_total = money(raw_daily)
                        continue
                    quantity = money(raw_quantity)
                    unit_price = money(raw_price)
                    amount = money(raw_amount)
                    if quantity is None or quantity <= 0:
                        result.issue('invalid_quantity', current.workbook, current.sheet, row_number, product)
                    else:
                        if amount is None and unit_price is not None:
                            amount = (quantity * unit_price).quantize(CENT, rounding=ROUND_HALF_UP)
                            result.issue('calculated_line_amount', current.workbook, current.sheet, row_number, product)
                        if unit_price is None and current.lines and all(line.amount is None for line in current.lines):
                            amount = None
                        if amount is None:
                            result.issue('missing_line_amount', current.workbook, current.sheet, row_number, product)
                        current.lines.append(SourceLine(row_number, product, quantity, unit_price, amount))
                if isinstance(raw_daily, (int, float, Decimal)):
                    current.reported_total = money(raw_daily)
                for extra in padded[7:]:
                    if isinstance(extra, str) and extra.strip():
                        current.notes.append(extra.strip())
            daily_sums = defaultdict(lambda: Decimal('0.00'))
            for ticket in sheet_tickets:
                if ticket.sale_date:
                    daily_sums[ticket.sale_date] += ticket.amount
            for ticket in sheet_tickets:
                if ticket.reported_total is None:
                    continue
                difference = ticket.amount - ticket.reported_total
                if difference and ticket.sale_date and abs(daily_sums[ticket.sale_date] - ticket.reported_total) <= CENT:
                    continue
                if difference:
                    result.daily_checks.append({'source': ticket.reference, 'reported': str(ticket.reported_total), 'parsed': str(ticket.amount), 'difference': str(difference)})
            if monthly_reported is not None:
                parsed = sum(daily_sums.values(), Decimal('0.00'))
                result.monthly_checks.append({'source': f'{workbook_year}:{sheet.title.strip()}', 'reported': str(monthly_reported), 'parsed': str(parsed), 'difference': str(parsed - monthly_reported)})
        workbook.close()
    by_invoice = defaultdict(list)
    for ticket in result.tickets:
        if ticket.invoice_number:
            by_invoice[(ticket.workbook, ticket.invoice_number, canonical_customer(ticket.customer_label))].append(ticket)
    for group in by_invoice.values():
        may = next((ticket for ticket in group if ticket.workbook == '2024' and ticket.sheet == 'MAY'), None)
        if may:
            for ticket in group:
                if ticket.sheet == 'JANUARY' and ticket is not may:
                    ticket.duplicate_of = may.reference
                    result.issue('duplicate_cross_sheet', ticket.workbook, ticket.sheet, ticket.row, may.reference)
    return result
